from __future__ import annotations

import io
import json
import math
import re
import statistics
import uuid
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import pandas as pd
import pymupdf as fitz
import streamlit as st
from PIL import Image, ImageOps, UnidentifiedImageError


PDF_OCR_DPI = 300
NATIVE_TEXT_MIN_PRINTABLE_CHARS = 50
NATIVE_PRINTABLE_RATIO_THRESHOLD = 0.9
IMAGE_HEAVY_COVERAGE_THRESHOLD = 0.55

INSTALL_COMMANDS = """python -m pip install --upgrade pip
python -m pip install streamlit pymupdf easyocr opencv-python-headless pillow pandas numpy openpyxl xlrd pyxlsb odfpy python-docx python-pptx
streamlit run ai_studio_code.py
"""

SUPPORTED_UPLOAD_TYPES = [
    "pdf",
    "png",
    "jpg",
    "jpeg",
    "tif",
    "tiff",
    "bmp",
    "webp",
    "gif",
    "csv",
    "tsv",
    "xlsx",
    "xls",
    "xlsm",
    "xlsb",
    "ods",
    "docx",
    "pptx",
    "txt",
    "md",
    "log",
    "json",
]

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp", "gif"}
SPREADSHEET_EXTENSIONS = {"xlsx", "xls", "xlsm", "xlsb", "ods"}
TEXT_EXTENSIONS = {"txt", "md", "log", "json"}

DOMAIN_CONFUSIONS = [
    (r"\bASTM\s+A1O6\b", "ASTM A106", "OCR 'O' substituted for digit '0' in ASTM grade"),
    (r"\bASTM\s+A1O5\b", "ASTM A105", "OCR 'O' substituted for digit '0' in ASTM grade"),
    (r"\bSCH4O\b", "SCH40", "OCR 'O' substituted for digit '0' in pipe schedule"),
    (r"\bSCH8O\b", "SCH80", "OCR 'O' substituted for digit '0' in pipe schedule"),
    (r"\bSS3I6\b", "SS316", "OCR 'I' substituted for digit '1' in stainless grade"),
    (r"\bSS3O4\b", "SS304", "OCR 'O' substituted for digit '0' in stainless grade"),
    (r"\b49O5(\d+)\b", r"4905\1", "OCR 'O' substituted for digit '0' in material code"),
]

HEADER_TOKENS = {
    "sis",
    "code",
    "codes",
    "group",
    "sap",
    "description",
    "descriptions",
    "uom",
    "unit",
    "qty",
    "quantity",
    "item",
    "material",
}


@dataclass
class CandidateCorrection:
    candidate_value: str
    reason: str
    rule: str = "domain_digit_letter_confusion"
    confidence: float = 0.95


@dataclass
class TextElement:
    id: str
    page: int
    raw_text: str
    bbox: list[float]
    confidence: float
    engine: str
    route: str
    source_document: str
    coordinate_space: str
    provenance: dict[str, Any]
    candidate_corrections: list[CandidateCorrection] = field(default_factory=list)


@dataclass
class TableCell:
    row: int
    column: int
    text: str
    bbox: list[float]
    confidence: float
    source_element_ids: list[str]
    provenance: dict[str, Any]


@dataclass
class ReconstructedTable:
    id: str
    page: int
    source_document: str
    engine: str
    route: str
    coordinate_space: str
    bbox: list[float]
    confidence: float
    headers: list[str]
    rows: list[list[str]]
    cells: list[TableCell]
    reconstruction_method: str


@dataclass
class PageResult:
    page: int
    route: str
    reason: str
    native_printable_chars: int
    native_printable_ratio: float
    image_coverage: float
    render_dpi: int | None
    elements: list[TextElement]
    tables: list[ReconstructedTable]
    warnings: list[str] = field(default_factory=list)
    preview_image: Image.Image | None = None


def printable_text_stats(text: str) -> tuple[int, float]:
    if not text:
        return 0, 0.0
    printable = sum(1 for char in text if char.isprintable() and not char.isspace())
    non_space = sum(1 for char in text if not char.isspace())
    ratio = printable / non_space if non_space else 0.0
    return printable, ratio


def validate_domain_text(raw_text: str) -> list[CandidateCorrection]:
    corrections: list[CandidateCorrection] = []
    for pattern, replacement, reason in DOMAIN_CONFUSIONS:
        if re.search(pattern, raw_text, re.IGNORECASE):
            candidate = re.sub(pattern, replacement, raw_text, flags=re.IGNORECASE)
            corrections.append(CandidateCorrection(candidate_value=candidate, reason=reason))
    return corrections


def clean_text(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def bbox_from_points(points: Any) -> list[float]:
    xs: list[float] = []
    ys: list[float] = []
    for point in points:
        xs.append(float(point[0]))
        ys.append(float(point[1]))
    if not xs or not ys:
        return [0.0, 0.0, 0.0, 0.0]
    return [round(min(xs), 2), round(min(ys), 2), round(max(xs), 2), round(max(ys), 2)]


def union_bbox(boxes: list[list[float]]) -> list[float]:
    valid = [box for box in boxes if len(box) == 4]
    if not valid:
        return [0.0, 0.0, 0.0, 0.0]
    return [
        round(min(box[0] for box in valid), 2),
        round(min(box[1] for box in valid), 2),
        round(max(box[2] for box in valid), 2),
        round(max(box[3] for box in valid), 2),
    ]


def bbox_center_y(element: TextElement) -> float:
    return (element.bbox[1] + element.bbox[3]) / 2


def bbox_start_x(element: TextElement) -> float:
    return element.bbox[0]


def bbox_width(element: TextElement) -> float:
    return max(1.0, element.bbox[2] - element.bbox[0])


def bbox_height(element: TextElement) -> float:
    return max(1.0, element.bbox[3] - element.bbox[1])


def make_provenance(
    page: int,
    source_document: str,
    engine: str,
    route: str,
    coordinate_space: str,
    **extra: Any,
) -> dict[str, Any]:
    provenance = {
        "page": page,
        "source_document": source_document,
        "engine": engine,
        "route": route,
        "coordinate_space": coordinate_space,
    }
    provenance.update(extra)
    return provenance


def dataclass_to_plain(value: Any) -> Any:
    if hasattr(value, "__dataclass_fields__"):
        return {key: dataclass_to_plain(item) for key, item in asdict(value).items() if key != "preview_image"}
    if isinstance(value, list):
        return [dataclass_to_plain(item) for item in value]
    if isinstance(value, dict):
        return {key: dataclass_to_plain(item) for key, item in value.items()}
    return value


@st.cache_resource(show_spinner="Loading EasyOCR model. First run may download model files.")
def get_easyocr_reader(languages: tuple[str, ...], use_gpu: bool):
    try:
        import easyocr
    except ImportError as exc:
        raise RuntimeError(
            "EasyOCR is not installed. Run: python -m pip install easyocr"
        ) from exc

    try:
        return easyocr.Reader(list(languages), gpu=use_gpu, verbose=False)
    except Exception as exc:
        model_directory = Path.home() / ".EasyOCR" / "model"
        raise RuntimeError(
            "EasyOCR could not initialize. Run the install commands in the sidebar, "
            "then retry. On first use EasyOCR downloads model files into "
            f"`{model_directory}`; if Python's downloader is blocked, "
            "download `craft_mlt_25k.zip` and `english_g2.zip` from the EasyOCR GitHub "
            "release assets and extract the `.pth` files into that model folder."
        ) from exc


def preprocess_for_ocr(image: Image.Image) -> tuple[np.ndarray, dict[str, Any]]:
    oriented_image = ImageOps.exif_transpose(image)
    rgb = np.array(oriented_image.convert("RGB"))
    height, width = rgb.shape[:2]
    scale = 1.0

    target_min_dimension = 1600
    target_max_dimension = 2880
    longest_dimension = max(width, height)
    if longest_dimension < target_min_dimension:
        scale = target_min_dimension / longest_dimension
    elif longest_dimension > target_max_dimension:
        scale = target_max_dimension / longest_dimension

    if scale != 1.0:
        interpolation = cv2.INTER_CUBIC if scale > 1.0 else cv2.INTER_AREA
        rgb = cv2.resize(rgb, None, fx=scale, fy=scale, interpolation=interpolation)

    return rgb, {
        "original_width": width,
        "original_height": height,
        "processed_width": int(rgb.shape[1]),
        "processed_height": int(rgb.shape[0]),
        "scale": round(scale, 4),
        "steps": ["exif_orientation", "rgb_convert", "bounded_resize"],
    }


def enhance_ocr_image(image: np.ndarray) -> np.ndarray:
    gray = cv2.cvtColor(image, cv2.COLOR_RGB2GRAY)
    denoised = cv2.medianBlur(gray, 3)
    clahe = cv2.createCLAHE(clipLimit=2.2, tileGridSize=(8, 8))
    contrast = clahe.apply(denoised)
    sharpened = cv2.addWeighted(contrast, 1.5, cv2.GaussianBlur(contrast, (0, 0), 1.0), -0.5, 0)
    return cv2.cvtColor(sharpened, cv2.COLOR_GRAY2RGB)


def read_easyocr_pass(reader: Any, image: np.ndarray) -> list[Any]:
    return reader.readtext(
        image,
        detail=1,
        paragraph=False,
        batch_size=4,
        workers=0,
        canvas_size=2048,
        mag_ratio=1.0,
        min_size=5,
        text_threshold=0.6,
        low_text=0.3,
        link_threshold=0.3,
        slope_ths=0.15,
        ycenter_ths=0.5,
        height_ths=0.5,
        width_ths=0.75,
    )


def detect_sideways_text(image: np.ndarray) -> dict[str, Any]:
    gray = cv2.cvtColor(image, cv2.COLOR_RGB2GRAY)
    height, width = gray.shape[:2]
    longest_side = max(width, height)
    if longest_side > 1600:
        scale = 1600.0 / longest_side
        gray = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
        height, width = gray.shape[:2]

    binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    horizontal_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (max(20, width // 35), 1))
    vertical_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(20, height // 35)))
    horizontal_lines = cv2.morphologyEx(binary, cv2.MORPH_OPEN, horizontal_kernel)
    vertical_lines = cv2.morphologyEx(binary, cv2.MORPH_OPEN, vertical_kernel)
    line_mask = cv2.bitwise_or(horizontal_lines, vertical_lines)
    text_mask = cv2.bitwise_and(binary, cv2.bitwise_not(line_mask))

    foreground_pixels = int(cv2.countNonZero(text_mask))
    if foreground_pixels < max(100, int(width * height * 0.0001)):
        text_mask = binary
        foreground_pixels = int(cv2.countNonZero(text_mask))

    row_density = np.count_nonzero(text_mask, axis=1) / max(1, width)
    column_density = np.count_nonzero(text_mask, axis=0) / max(1, height)
    row_variation = float(row_density.std() / (row_density.mean() + 1e-9))
    column_variation = float(column_density.std() / (column_density.mean() + 1e-9))
    sideways = foreground_pixels > 0 and column_variation > row_variation * 1.35
    return {
        "sideways": sideways,
        "row_projection_variation": round(row_variation, 4),
        "column_projection_variation": round(column_variation, 4),
        "foreground_pixels": foreground_pixels,
    }


def summarize_ocr_quality(results: list[Any]) -> dict[str, Any]:
    texts: list[str] = []
    confidences: list[float] = []
    alphanumeric_lengths: list[int] = []
    for result in results:
        if len(result) < 3:
            continue
        text = clean_text(result[1])
        if not text:
            continue
        texts.append(text)
        confidences.append(max(0.0, min(1.0, float(result[2]))))
        alphanumeric_lengths.append(len(re.sub(r"[^A-Za-z0-9]", "", text)))

    count = len(texts)
    character_count = sum(alphanumeric_lengths)
    single_character_count = sum(length <= 1 for length in alphanumeric_lengths)
    single_character_ratio = single_character_count / count if count else 0.0
    mean_confidence = statistics.mean(confidences) if confidences else 0.0
    meaningful_tokens = sum(length >= 2 for length in alphanumeric_lengths)
    score = character_count * (0.5 + mean_confidence) + meaningful_tokens * 4 - single_character_count * 2
    return {
        "result_count": count,
        "character_count": character_count,
        "single_character_ratio": round(single_character_ratio, 4),
        "mean_confidence": round(mean_confidence, 4),
        "score": round(score, 4),
    }


def needs_orientation_retry(quality: dict[str, Any]) -> bool:
    result_count = int(quality["result_count"])
    return (
        result_count >= 4
        and float(quality["single_character_ratio"]) >= 0.65
        and int(quality["character_count"]) <= result_count * 2
    )


def needs_enhancement_retry(quality: dict[str, Any]) -> bool:
    return (
        int(quality["result_count"]) == 0
        or int(quality["character_count"]) < 8
        or float(quality["mean_confidence"]) < 0.45
    )


def rotate_ocr_image(image: np.ndarray, degrees_clockwise: int) -> np.ndarray:
    if degrees_clockwise == 90:
        return cv2.rotate(image, cv2.ROTATE_90_CLOCKWISE)
    if degrees_clockwise == 180:
        return cv2.rotate(image, cv2.ROTATE_180)
    if degrees_clockwise == 270:
        return cv2.rotate(image, cv2.ROTATE_90_COUNTERCLOCKWISE)
    return image


def run_easyocr(
    image: Image.Image,
    page_number: int,
    source_document: str,
    route: str,
    coordinate_space: str,
    languages: tuple[str, ...],
    use_gpu: bool,
) -> tuple[list[TextElement], dict[str, Any]]:
    base_processed, metadata = preprocess_for_ocr(image)
    reader = get_easyocr_reader(languages, use_gpu)

    orientation_hint = detect_sideways_text(base_processed)
    selected_rotation = 90 if bool(orientation_hint["sideways"]) else 0
    selected_variant = "bounded_rgb"
    processed = rotate_ocr_image(base_processed, selected_rotation)
    results = read_easyocr_pass(reader, processed)
    quality = summarize_ocr_quality(results)
    attempts = [{"rotation_degrees_clockwise": selected_rotation, **quality}]

    if needs_orientation_retry(quality):
        best_score = float(quality["score"])
        rotations = (270, 0, 180) if selected_rotation == 90 else (90, 270, 180)
        for rotation in rotations:
            rotated = rotate_ocr_image(base_processed, rotation)
            candidate_results = read_easyocr_pass(reader, rotated)
            candidate_quality = summarize_ocr_quality(candidate_results)
            attempts.append({"rotation_degrees_clockwise": rotation, **candidate_quality})
            if float(candidate_quality["score"]) > best_score:
                processed = rotated
                results = candidate_results
                quality = candidate_quality
                best_score = float(candidate_quality["score"])
                selected_rotation = rotation
            if (
                int(candidate_quality["character_count"]) >= 20
                and float(candidate_quality["single_character_ratio"]) < 0.4
            ):
                break

    if needs_enhancement_retry(quality):
        enhanced = rotate_ocr_image(enhance_ocr_image(base_processed), selected_rotation)
        enhanced_results = read_easyocr_pass(reader, enhanced)
        enhanced_quality = summarize_ocr_quality(enhanced_results)
        attempts.append(
            {
                "rotation_degrees_clockwise": selected_rotation,
                "image_variant": "contrast_enhanced",
                **enhanced_quality,
            }
        )
        if float(enhanced_quality["score"]) > float(quality["score"]):
            processed = enhanced
            results = enhanced_results
            quality = enhanced_quality
            selected_variant = "contrast_enhanced"

    metadata["processed_width"] = int(processed.shape[1])
    metadata["processed_height"] = int(processed.shape[0])
    metadata["ocr_rotation_degrees_clockwise"] = selected_rotation
    metadata["orientation_projection"] = orientation_hint
    metadata["ocr_image_variant"] = selected_variant
    metadata["ocr_quality"] = quality
    metadata["ocr_attempts"] = attempts
    if len(attempts) > 1:
        metadata["steps"].append("quality_retry")

    elements: list[TextElement] = []
    for result_index, result in enumerate(results, start=1):
        if len(result) < 3:
            continue
        points, text, confidence = result[0], clean_text(result[1]), float(result[2])
        if not text:
            continue

        bbox = bbox_from_points(points)
        provenance = make_provenance(
            page_number,
            source_document,
            "easyocr",
            route,
            coordinate_space,
            ocr_result_index=result_index,
            preprocessing=metadata,
        )
        elements.append(
            TextElement(
                id=f"OCR-{page_number}-{uuid.uuid4().hex[:8]}",
                page=page_number,
                raw_text=text,
                bbox=bbox,
                confidence=round(confidence, 4),
                engine="easyocr",
                route=route,
                source_document=source_document,
                coordinate_space=coordinate_space,
                provenance=provenance,
                candidate_corrections=validate_domain_text(text),
            )
        )

    elements.sort(key=lambda item: (bbox_center_y(item), bbox_start_x(item)))
    return elements, metadata


def estimate_image_coverage(page: fitz.Page) -> float:
    page_area = max(1.0, float(page.rect.width * page.rect.height))
    image_area = 0.0

    try:
        for info in page.get_image_info():
            bbox = info.get("bbox")
            if not bbox:
                continue
            x0, y0, x1, y1 = [float(value) for value in bbox]
            image_area += max(0.0, x1 - x0) * max(0.0, y1 - y0)
    except Exception:
        image_area = 0.0

    if image_area == 0.0:
        try:
            for image_info in page.get_images(full=True):
                xref = image_info[0]
                for rect in page.get_image_rects(xref):
                    image_area += max(0.0, rect.width * rect.height)
        except Exception:
            image_area = 0.0

    return round(min(1.0, image_area / page_area), 4)


def classify_pdf_page(page: fitz.Page) -> dict[str, Any]:
    native_text = page.get_text("text") or ""
    printable_chars, printable_ratio = printable_text_stats(native_text)
    image_coverage = estimate_image_coverage(page)

    has_clean_native_text = (
        printable_chars > NATIVE_TEXT_MIN_PRINTABLE_CHARS
        and printable_ratio >= NATIVE_PRINTABLE_RATIO_THRESHOLD
    )
    if has_clean_native_text:
        return {
            "route": "native_pdf",
            "reason": "clean native text exceeds threshold",
            "native_printable_chars": printable_chars,
            "native_printable_ratio": printable_ratio,
            "image_coverage": image_coverage,
        }

    if printable_chars == 0:
        reason = "no native text"
    elif image_coverage >= IMAGE_HEAVY_COVERAGE_THRESHOLD:
        reason = "image-heavy page with weak native text"
    else:
        reason = "native text below clean-text threshold"

    return {
        "route": "ocr_pdf_scan",
        "reason": reason,
        "native_printable_chars": printable_chars,
        "native_printable_ratio": printable_ratio,
        "image_coverage": image_coverage,
    }


def render_pdf_page(page: fitz.Page, dpi: int) -> Image.Image:
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    return Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")


def extract_native_pdf_elements(
    page: fitz.Page,
    page_number: int,
    source_document: str,
) -> list[TextElement]:
    elements: list[TextElement] = []
    for block_index, block in enumerate(page.get_text("blocks", sort=True), start=1):
        if len(block) < 5:
            continue
        raw_text = clean_text(block[4])
        if not raw_text:
            continue
        bbox = [round(float(value), 2) for value in block[:4]]
        provenance = make_provenance(
            page_number,
            source_document,
            "pymupdf",
            "native_pdf",
            "pdf_points",
            block_index=block_index,
        )
        elements.append(
            TextElement(
                id=f"NATIVE-{page_number}-{uuid.uuid4().hex[:8]}",
                page=page_number,
                raw_text=raw_text,
                bbox=bbox,
                confidence=1.0,
                engine="pymupdf",
                route="native_pdf",
                source_document=source_document,
                coordinate_space="pdf_points",
                provenance=provenance,
                candidate_corrections=validate_domain_text(raw_text),
            )
        )
    return elements


def extract_native_pdf_tables(
    page: fitz.Page,
    page_number: int,
    source_document: str,
) -> tuple[list[ReconstructedTable], list[str]]:
    warnings: list[str] = []
    tables: list[ReconstructedTable] = []

    if not hasattr(page, "find_tables"):
        return tables, ["PyMuPDF table extraction is unavailable in this installation."]

    try:
        table_finder = page.find_tables()
    except Exception as exc:
        return tables, [f"Native table extraction failed on page {page_number}: {exc}"]

    for table_index, table in enumerate(getattr(table_finder, "tables", []), start=1):
        raw_rows = table.extract() or []
        rows = [[clean_text(cell) for cell in row] for row in raw_rows]
        if not rows:
            continue

        header_names = [clean_text(name) for name in getattr(getattr(table, "header", None), "names", [])]
        headers = header_names if any(header_names) else [f"Column {index + 1}" for index in range(max(len(row) for row in rows))]

        bbox = [round(float(value), 2) for value in getattr(table, "bbox", [0.0, 0.0, 0.0, 0.0])]
        cells: list[TableCell] = []
        pymupdf_cells = list(getattr(table, "cells", []) or [])
        flat_index = 0
        for row_index, row in enumerate(rows):
            for col_index, text in enumerate(row):
                cell_bbox = [0.0, 0.0, 0.0, 0.0]
                if flat_index < len(pymupdf_cells) and pymupdf_cells[flat_index]:
                    cell_bbox = [round(float(value), 2) for value in pymupdf_cells[flat_index]]
                flat_index += 1
                cells.append(
                    TableCell(
                        row=row_index,
                        column=col_index,
                        text=text,
                        bbox=cell_bbox,
                        confidence=1.0,
                        source_element_ids=[],
                        provenance=make_provenance(
                            page_number,
                            source_document,
                            "pymupdf-table",
                            "native_pdf",
                            "pdf_points",
                            table_index=table_index,
                            row=row_index,
                            column=col_index,
                        ),
                    )
                )

        tables.append(
            ReconstructedTable(
                id=f"TABLE-{page_number}-{table_index}",
                page=page_number,
                source_document=source_document,
                engine="pymupdf-table",
                route="native_pdf",
                coordinate_space="pdf_points",
                bbox=bbox,
                confidence=1.0,
                headers=headers,
                rows=rows,
                cells=cells,
                reconstruction_method="pymupdf_native_table",
            )
        )

    return tables, warnings


def group_elements_into_rows(elements: list[TextElement]) -> list[list[TextElement]]:
    if not elements:
        return []

    heights = [bbox_height(element) for element in elements]
    y_tolerance = max(12.0, statistics.median(heights) * 0.75)

    rows: list[list[TextElement]] = []
    row_centers: list[float] = []
    for element in sorted(elements, key=lambda item: (bbox_center_y(item), bbox_start_x(item))):
        center_y = bbox_center_y(element)
        best_index = None
        best_distance = math.inf
        for row_index, row_center in enumerate(row_centers):
            distance = abs(center_y - row_center)
            if distance <= y_tolerance and distance < best_distance:
                best_index = row_index
                best_distance = distance

        if best_index is None:
            rows.append([element])
            row_centers.append(center_y)
        else:
            rows[best_index].append(element)
            row_centers[best_index] = statistics.mean([bbox_center_y(item) for item in rows[best_index]])

    for row in rows:
        row.sort(key=bbox_start_x)

    return sorted(rows, key=lambda row: statistics.mean([bbox_center_y(item) for item in row]))


def infer_column_starts(rows: list[list[TextElement]], image_width: int) -> list[float]:
    x_values = sorted(bbox_start_x(element) for row in rows for element in row)
    if not x_values:
        return []

    widths = [bbox_width(element) for row in rows for element in row]
    median_width = statistics.median(widths) if widths else 40.0
    threshold = max(28.0, min(float(image_width) * 0.045, median_width * 0.85))

    clusters: list[list[float]] = []
    for x_value in x_values:
        if not clusters or abs(x_value - statistics.median(clusters[-1])) > threshold:
            clusters.append([x_value])
        else:
            clusters[-1].append(x_value)

    starts = [round(float(statistics.median(cluster)), 2) for cluster in clusters]

    deduped: list[float] = []
    for start in starts:
        if not deduped or abs(start - deduped[-1]) > 12:
            deduped.append(start)
    return deduped


def nearest_column_index(x_value: float, column_starts: list[float]) -> int:
    if not column_starts:
        return 0
    return min(range(len(column_starts)), key=lambda index: abs(x_value - column_starts[index]))


def is_probable_header(row: list[str]) -> bool:
    joined = " ".join(row).lower()
    tokens = set(re.findall(r"[a-z0-9]+", joined))
    return len(tokens & HEADER_TOKENS) >= 2


def normalize_headers(headers: list[str], column_count: int) -> list[str]:
    normalized: list[str] = []
    seen: dict[str, int] = {}
    for index in range(column_count):
        header = clean_text(headers[index] if index < len(headers) else "") or f"Column {index + 1}"
        key = header.lower()
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1:
            header = f"{header} {seen[key]}"
        normalized.append(header)
    return normalized


def split_known_header_labels(header: str) -> list[str]:
    text = clean_text(header)
    if not text:
        return []

    lowered = text.lower()
    known_labels = [
        ("sis code", "SIS Code"),
        ("sap code", "SAP Code"),
        ("material code", "Material Code"),
        ("description", "Description"),
        ("group", "Group"),
        ("uom", "UOM"),
        ("unit", "Unit"),
        ("quantity", "Quantity"),
        ("qty", "Qty"),
        ("item", "Item"),
        ("code", "Code"),
    ]

    found: list[tuple[int, str]] = []
    specific_code_seen = False
    for needle, label in known_labels:
        if needle == "code" and specific_code_seen:
            continue
        match = re.search(rf"(?<![a-z0-9]){re.escape(needle)}(?![a-z0-9])", lowered)
        if match:
            if needle.endswith(" code") and needle != "code":
                specific_code_seen = True
            found.append((match.start(), label))

    if len(found) >= 2:
        return [label for _, label in sorted(found, key=lambda item: item[0])]
    return [text]


def reconstruct_spatial_table(
    elements: list[TextElement],
    page_number: int,
    source_document: str,
    route: str,
    coordinate_space: str,
    image_width: int,
) -> list[ReconstructedTable]:
    rows = group_elements_into_rows(elements)
    if not rows:
        return []

    column_starts = infer_column_starts(rows, image_width)
    column_count = max(1, len(column_starts))

    row_values: list[list[str]] = []
    row_groups: list[list[list[TextElement]]] = []

    for row in rows:
        grouped: list[list[TextElement]] = [[] for _ in range(column_count)]
        for element in row:
            col_index = nearest_column_index(bbox_start_x(element), column_starts)
            grouped[col_index].append(element)

        values: list[str] = []
        for cell_elements in grouped:
            cell_elements.sort(key=bbox_start_x)
            text = clean_text(" ".join(element.raw_text for element in cell_elements))
            values.append(text)
        row_values.append(values)
        row_groups.append(grouped)

    non_empty = [
        (row_index, row, row_groups[row_index])
        for row_index, row in enumerate(row_values)
        if any(cell for cell in row)
    ]
    if not non_empty:
        return []

    first_row = non_empty[0][1]
    header_offset = 1 if is_probable_header(first_row) else 0
    data_items = non_empty[header_offset:]
    header_labels = [""] * column_count

    if header_offset:
        occupied_columns = sorted(
            {
                col_index
                for _, row, _ in data_items
                for col_index, value in enumerate(row)
                if value
            }
        )
        if not occupied_columns:
            occupied_columns = sorted(
                {
                    col_index
                    for _, row, _ in non_empty
                    for col_index, value in enumerate(row)
                    if value
                }
            )

        for col_index, header in enumerate(first_row):
            header = clean_text(header)
            if not header:
                continue

            if col_index in occupied_columns or not occupied_columns:
                target = col_index
            else:
                target = min(
                    occupied_columns,
                    key=lambda occupied: abs(column_starts[occupied] - column_starts[col_index]),
                )

            header_labels[target] = clean_text(" ".join(part for part in [header_labels[target], header] if part))

        for col_index in occupied_columns:
            parts = split_known_header_labels(header_labels[col_index])
            if len(parts) <= 1:
                continue

            header_labels[col_index] = parts[0]
            open_columns = [
                candidate
                for candidate in occupied_columns
                if candidate > col_index and not header_labels[candidate]
            ]
            for candidate, label in zip(open_columns, parts[1:]):
                header_labels[candidate] = label
    else:
        occupied_columns = sorted(
            {
                col_index
                for _, row, _ in non_empty
                for col_index, value in enumerate(row)
                if value
            }
        )

    keep_columns = occupied_columns or list(range(column_count))
    headers = normalize_headers([header_labels[col_index] for col_index in keep_columns], len(keep_columns))
    data_rows = [
        [row[col_index] if col_index < len(row) else "" for col_index in keep_columns]
        for _, row, _ in data_items
    ]
    row_has_multiple_columns = sum(1 for row in data_rows if sum(1 for value in row if value) > 1)

    table_cells: list[TableCell] = []
    for compact_row_index, (_, _, grouped) in enumerate(data_items):
        for compact_col_index, original_col_index in enumerate(keep_columns):
            cell_elements = grouped[original_col_index] if original_col_index < len(grouped) else []
            cell_elements.sort(key=bbox_start_x)
            text = clean_text(" ".join(element.raw_text for element in cell_elements))
            if not text:
                continue

            confidence = round(statistics.mean([element.confidence for element in cell_elements]), 4)
            table_cells.append(
                TableCell(
                    row=compact_row_index,
                    column=compact_col_index,
                    text=text,
                    bbox=union_bbox([element.bbox for element in cell_elements]),
                    confidence=confidence,
                    source_element_ids=[element.id for element in cell_elements],
                    provenance=make_provenance(
                        page_number,
                        source_document,
                        "easyocr-spatial-reconstruction",
                        route,
                        coordinate_space,
                        row=compact_row_index,
                        column=compact_col_index,
                    ),
                )
            )

    reconstruction_kind = (
        "ocr_yx_clustered_table"
        if len(keep_columns) > 1 and row_has_multiple_columns >= 2
        else "ocr_y_sorted_text_rows"
    )

    table_bbox = union_bbox([element.bbox for element in elements])
    confidence_values = [element.confidence for element in elements]
    table_confidence = round(statistics.mean(confidence_values), 4) if confidence_values else 0.0

    return [
        ReconstructedTable(
            id=f"TABLE-{page_number}-OCR",
            page=page_number,
            source_document=source_document,
            engine="easyocr",
            route=route,
            coordinate_space=coordinate_space,
            bbox=table_bbox,
            confidence=table_confidence,
            headers=headers,
            rows=data_rows,
            cells=table_cells,
            reconstruction_method=reconstruction_kind,
        )
    ]


def process_pdf(
    file_bytes: bytes,
    filename: str,
    languages: tuple[str, ...],
    use_gpu: bool,
) -> list[PageResult]:
    pages: list[PageResult] = []
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise RuntimeError(f"Could not open PDF: {exc}") from exc

    if doc.needs_pass:
        doc.close()
        raise RuntimeError("This PDF is password protected. Upload an unlocked copy.")

    for page_index, page in enumerate(doc, start=1):
        route_info = classify_pdf_page(page)
        route = route_info["route"]
        preview = render_pdf_page(page, dpi=150)
        warnings: list[str] = []

        if route == "native_pdf":
            elements = extract_native_pdf_elements(page, page_index, filename)
            tables, table_warnings = extract_native_pdf_tables(page, page_index, filename)
            warnings.extend(table_warnings)
            render_dpi = None
        else:
            ocr_image = render_pdf_page(page, dpi=PDF_OCR_DPI)
            coordinate_space = f"ocr_preprocessed_pixels_from_{PDF_OCR_DPI}dpi_pdf_render"
            elements, metadata = run_easyocr(
                ocr_image,
                page_index,
                filename,
                route,
                coordinate_space,
                languages,
                use_gpu,
            )
            tables = reconstruct_spatial_table(
                elements,
                page_index,
                filename,
                route,
                coordinate_space,
                int(metadata["processed_width"]),
            )
            if int(metadata["ocr_rotation_degrees_clockwise"]):
                warnings.append(
                    f"OCR auto-rotated this page {metadata['ocr_rotation_degrees_clockwise']} degrees clockwise."
                )
            if not elements:
                warnings.append("OCR completed but did not detect readable text on this page.")
            render_dpi = PDF_OCR_DPI

        pages.append(
            PageResult(
                page=page_index,
                route=route,
                reason=str(route_info["reason"]),
                native_printable_chars=int(route_info["native_printable_chars"]),
                native_printable_ratio=round(float(route_info["native_printable_ratio"]), 4),
                image_coverage=round(float(route_info["image_coverage"]), 4),
                render_dpi=render_dpi,
                elements=elements,
                tables=tables,
                warnings=warnings,
                preview_image=preview,
            )
        )

    doc.close()
    return pages


def process_image(
    file_bytes: bytes,
    filename: str,
    languages: tuple[str, ...],
    use_gpu: bool,
) -> list[PageResult]:
    try:
        source_image = Image.open(io.BytesIO(file_bytes))
    except UnidentifiedImageError as exc:
        raise RuntimeError("Could not open this image file.") from exc

    pages: list[PageResult] = []
    frame_count = int(getattr(source_image, "n_frames", 1))
    for frame_index in range(frame_count):
        source_image.seek(frame_index)
        image = ImageOps.exif_transpose(source_image.copy()).convert("RGB")
        page_number = frame_index + 1
        elements, metadata = run_easyocr(
            image,
            page_number,
            filename,
            "ocr_image",
            "ocr_preprocessed_image_pixels",
            languages,
            use_gpu,
        )
        tables = reconstruct_spatial_table(
            elements,
            page_number,
            filename,
            "ocr_image",
            "ocr_preprocessed_image_pixels",
            int(metadata["processed_width"]),
        )
        warnings: list[str] = []
        if int(metadata["ocr_rotation_degrees_clockwise"]):
            warnings.append(
                f"OCR auto-rotated this image {metadata['ocr_rotation_degrees_clockwise']} degrees clockwise."
            )
        if not elements:
            warnings.append("OCR completed but did not detect readable text in this image.")
        pages.append(
            PageResult(
                page=page_number,
                route="ocr_image",
                reason="uploaded raster image" if frame_count == 1 else "image frame routed to OCR",
                native_printable_chars=0,
                native_printable_ratio=0.0,
                image_coverage=1.0,
                render_dpi=None,
                elements=elements,
                tables=tables,
                warnings=warnings,
                preview_image=image,
            )
        )

    source_image.close()
    return pages


def raw_value_to_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        missing = pd.isna(value)
        if isinstance(missing, (bool, np.bool_)) and bool(missing):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value)


def make_native_element(
    element_id: str,
    page_number: int,
    raw_text: str,
    filename: str,
    engine: str,
    route: str,
    coordinate_space: str,
    bbox: list[float] | None = None,
    **provenance_details: Any,
) -> TextElement:
    return TextElement(
        id=element_id,
        page=page_number,
        raw_text=raw_text,
        bbox=bbox or [0.0, 0.0, 0.0, 0.0],
        confidence=1.0,
        engine=engine,
        route=route,
        source_document=filename,
        coordinate_space=coordinate_space,
        provenance=make_provenance(
            page_number,
            filename,
            engine,
            route,
            coordinate_space,
            **provenance_details,
        ),
        candidate_corrections=validate_domain_text(raw_text),
    )


def make_structured_table(
    headers: list[str],
    rows: list[list[str]],
    page_number: int,
    filename: str,
    engine: str,
    route: str,
    table_id: str,
    reconstruction_method: str,
    **provenance_details: Any,
) -> tuple[ReconstructedTable, list[TextElement]]:
    width = max(len(headers), max((len(row) for row in rows), default=0))
    normalized_headers = normalize_headers(headers, width)
    normalized_rows = [row + [""] * (width - len(row)) for row in rows]
    elements: list[TextElement] = []
    cells: list[TableCell] = []

    for column_index, raw_header in enumerate(headers):
        if not raw_header:
            continue
        elements.append(
            make_native_element(
                f"{table_id}-HEADER-{column_index + 1}",
                page_number,
                raw_header,
                filename,
                engine,
                route,
                "structured_cells",
                row=-1,
                column=column_index,
                is_header=True,
                **provenance_details,
            )
        )

    for row_index, row in enumerate(normalized_rows):
        for column_index, text in enumerate(row):
            source_ids: list[str] = []
            if text:
                element = make_native_element(
                    f"{table_id}-R{row_index + 1}-C{column_index + 1}",
                    page_number,
                    text,
                    filename,
                    engine,
                    route,
                    "structured_cells",
                    row=row_index,
                    column=column_index,
                    **provenance_details,
                )
                elements.append(element)
                source_ids.append(element.id)

            cells.append(
                TableCell(
                    row=row_index,
                    column=column_index,
                    text=text,
                    bbox=[0.0, 0.0, 0.0, 0.0],
                    confidence=1.0,
                    source_element_ids=source_ids,
                    provenance=make_provenance(
                        page_number,
                        filename,
                        engine,
                        route,
                        "structured_cells",
                        row=row_index,
                        column=column_index,
                        **provenance_details,
                    ),
                )
            )

    table = ReconstructedTable(
        id=table_id,
        page=page_number,
        source_document=filename,
        engine=engine,
        route=route,
        coordinate_space="structured_cells",
        bbox=[0.0, 0.0, 0.0, 0.0],
        confidence=1.0,
        headers=normalized_headers,
        rows=normalized_rows,
        cells=cells,
        reconstruction_method=reconstruction_method,
    )
    return table, elements


def dataframe_to_page(
    dataframe: pd.DataFrame,
    page_number: int,
    filename: str,
    engine: str,
    route: str,
    reason: str,
    table_id: str,
    reconstruction_method: str,
    **provenance_details: Any,
) -> PageResult:
    headers = [raw_value_to_text(column) for column in dataframe.columns]
    rows = [
        [raw_value_to_text(value) for value in row]
        for row in dataframe.itertuples(index=False, name=None)
    ]
    table, elements = make_structured_table(
        headers,
        rows,
        page_number,
        filename,
        engine,
        route,
        table_id,
        reconstruction_method,
        **provenance_details,
    )
    printable_chars = sum(printable_text_stats(element.raw_text)[0] for element in elements)
    warnings = ["This sheet is empty."] if dataframe.empty and not headers else []
    return PageResult(
        page=page_number,
        route=route,
        reason=reason,
        native_printable_chars=printable_chars,
        native_printable_ratio=1.0,
        image_coverage=0.0,
        render_dpi=None,
        elements=elements,
        tables=[table],
        warnings=warnings,
        preview_image=None,
    )


def process_delimited(file_bytes: bytes, filename: str, extension: str) -> list[PageResult]:
    delimiter = "\t" if extension == "tsv" else None
    dataframe: pd.DataFrame | None = None
    last_error: Exception | None = None

    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            dataframe = pd.read_csv(
                io.BytesIO(file_bytes),
                sep=delimiter,
                engine="python" if delimiter is None else "c",
                dtype=str,
                keep_default_na=False,
                encoding=encoding,
            )
            break
        except UnicodeError as exc:
            last_error = exc

    if dataframe is None:
        raise RuntimeError(f"Could not decode {extension.upper()} file: {last_error}")

    return [
        dataframe_to_page(
            dataframe,
            1,
            filename,
            "pandas",
            f"native_{extension}",
            f"uploaded structured {extension.upper()} file",
            f"TABLE-1-{extension.upper()}",
            f"{extension}_native_parser",
        )
    ]


def process_csv(file_bytes: bytes, filename: str) -> list[PageResult]:
    return process_delimited(file_bytes, filename, "csv")


def process_spreadsheet(file_bytes: bytes, filename: str, extension: str) -> list[PageResult]:
    try:
        workbook = pd.ExcelFile(io.BytesIO(file_bytes))
    except ImportError as exc:
        raise RuntimeError(
            f"The reader for .{extension} is not installed. Run the install commands in the sidebar."
        ) from exc
    except Exception as exc:
        raise RuntimeError(f"Could not open spreadsheet: {exc}") from exc

    pages: list[PageResult] = []
    for sheet_index, sheet_name in enumerate(workbook.sheet_names, start=1):
        try:
            dataframe = workbook.parse(sheet_name=sheet_name, dtype=object, keep_default_na=False)
        except Exception as exc:
            raise RuntimeError(f"Could not read sheet '{sheet_name}': {exc}") from exc

        pages.append(
            dataframe_to_page(
                dataframe,
                sheet_index,
                filename,
                "pandas",
                f"native_{extension}",
                f"native spreadsheet sheet: {sheet_name}",
                f"TABLE-{sheet_index}-{extension.upper()}",
                f"{extension}_native_parser",
                sheet_name=sheet_name,
                sheet_index=sheet_index - 1,
            )
        )

    workbook.close()
    if not pages:
        raise RuntimeError("The spreadsheet contains no readable sheets.")
    return pages


def process_docx(file_bytes: bytes, filename: str) -> list[PageResult]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX support is not installed. Run the install commands in the sidebar.") from exc

    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise RuntimeError(f"Could not open DOCX file: {exc}") from exc

    route = "native_docx"
    elements: list[TextElement] = []
    tables: list[ReconstructedTable] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        raw_text = paragraph.text
        if not raw_text:
            continue
        elements.append(
            make_native_element(
                f"DOCX-P-{paragraph_index + 1}",
                1,
                raw_text,
                filename,
                "python-docx",
                route,
                "document_flow",
                paragraph_index=paragraph_index,
            )
        )

    for table_index, source_table in enumerate(document.tables, start=1):
        source_rows = [[cell.text for cell in row.cells] for row in source_table.rows]
        headers = source_rows[0] if source_rows else []
        rows = source_rows[1:] if source_rows else []
        table, table_elements = make_structured_table(
            headers,
            rows,
            1,
            filename,
            "python-docx",
            route,
            f"TABLE-1-DOCX-{table_index}",
            "docx_native_table",
            table_index=table_index - 1,
        )
        tables.append(table)
        elements.extend(table_elements)

    printable_chars = sum(printable_text_stats(element.raw_text)[0] for element in elements)
    return [
        PageResult(
            page=1,
            route=route,
            reason="native Word paragraphs and tables",
            native_printable_chars=printable_chars,
            native_printable_ratio=1.0,
            image_coverage=0.0,
            render_dpi=None,
            elements=elements,
            tables=tables,
            preview_image=None,
        )
    ]


def process_pptx(file_bytes: bytes, filename: str) -> list[PageResult]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX support is not installed. Run the install commands in the sidebar.") from exc

    try:
        presentation = Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise RuntimeError(f"Could not open PPTX file: {exc}") from exc

    route = "native_pptx"
    pages: list[PageResult] = []
    emu_per_point = 12700.0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        elements: list[TextElement] = []
        tables: list[ReconstructedTable] = []
        for shape_index, shape in enumerate(slide.shapes, start=1):
            shape_bbox = [
                round(float(shape.left) / emu_per_point, 2),
                round(float(shape.top) / emu_per_point, 2),
                round(float(shape.left + shape.width) / emu_per_point, 2),
                round(float(shape.top + shape.height) / emu_per_point, 2),
            ]

            if getattr(shape, "has_table", False):
                source_rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                headers = source_rows[0] if source_rows else []
                rows = source_rows[1:] if source_rows else []
                table, table_elements = make_structured_table(
                    headers,
                    rows,
                    slide_index,
                    filename,
                    "python-pptx",
                    route,
                    f"TABLE-{slide_index}-PPTX-{shape_index}",
                    "pptx_native_table",
                    slide_index=slide_index - 1,
                    shape_index=shape_index - 1,
                    shape_bbox_points=shape_bbox,
                )
                table.bbox = shape_bbox
                tables.append(table)
                elements.extend(table_elements)
                continue

            raw_text = getattr(shape, "text", "")
            if raw_text:
                elements.append(
                    make_native_element(
                        f"PPTX-{slide_index}-SHAPE-{shape_index}",
                        slide_index,
                        raw_text,
                        filename,
                        "python-pptx",
                        route,
                        "pptx_points",
                        bbox=shape_bbox,
                        slide_index=slide_index - 1,
                        shape_index=shape_index - 1,
                    )
                )

        printable_chars = sum(printable_text_stats(element.raw_text)[0] for element in elements)
        pages.append(
            PageResult(
                page=slide_index,
                route=route,
                reason="native PowerPoint slide text and tables",
                native_printable_chars=printable_chars,
                native_printable_ratio=1.0,
                image_coverage=0.0,
                render_dpi=None,
                elements=elements,
                tables=tables,
                preview_image=None,
            )
        )

    if not pages:
        raise RuntimeError("The presentation contains no slides.")
    return pages


def process_text_file(file_bytes: bytes, filename: str, extension: str) -> list[PageResult]:
    text: str | None = None
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = file_bytes.decode(encoding)
            break
        except UnicodeError as exc:
            last_error = exc
    if text is None:
        raise RuntimeError(f"Could not decode text file: {last_error}")

    route = f"native_{extension}"
    elements: list[TextElement] = []
    for line_index, line in enumerate(text.splitlines()):
        if not line:
            continue
        elements.append(
            make_native_element(
                f"TEXT-1-LINE-{line_index + 1}",
                1,
                line,
                filename,
                "python-text-decoder",
                route,
                "text_lines",
                line_number=line_index + 1,
            )
        )

    printable_chars, printable_ratio = printable_text_stats(text)
    return [
        PageResult(
            page=1,
            route=route,
            reason=f"native {extension.upper()} text",
            native_printable_chars=printable_chars,
            native_printable_ratio=printable_ratio,
            image_coverage=0.0,
            render_dpi=None,
            elements=elements,
            tables=[],
            preview_image=None,
        )
    ]


def process_uploaded_file(
    uploaded_file: Any,
    languages: tuple[str, ...],
    use_gpu: bool,
) -> tuple[list[PageResult], bytes, str]:
    file_bytes = uploaded_file.getvalue()
    filename = uploaded_file.name or "uploaded-file"
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if extension == "pdf":
        return process_pdf(file_bytes, filename, languages, use_gpu), file_bytes, filename
    if extension in IMAGE_EXTENSIONS:
        return process_image(file_bytes, filename, languages, use_gpu), file_bytes, filename
    if extension in {"csv", "tsv"}:
        return process_delimited(file_bytes, filename, extension), file_bytes, filename
    if extension in SPREADSHEET_EXTENSIONS:
        return process_spreadsheet(file_bytes, filename, extension), file_bytes, filename
    if extension == "docx":
        return process_docx(file_bytes, filename), file_bytes, filename
    if extension == "pptx":
        return process_pptx(file_bytes, filename), file_bytes, filename
    if extension in TEXT_EXTENSIONS:
        return process_text_file(file_bytes, filename, extension), file_bytes, filename

    supported = ", ".join(f".{item}" for item in SUPPORTED_UPLOAD_TYPES)
    raise RuntimeError(f"Unsupported file type: .{extension}. Supported types: {supported}")


def pages_to_json(pages: list[PageResult]) -> str:
    return json.dumps(dataclass_to_plain(pages), indent=2, ensure_ascii=False)


def elements_dataframe(pages: list[PageResult]) -> pd.DataFrame:
    records: list[dict[str, Any]] = []
    for page in pages:
        for element in page.elements:
            records.append(
                {
                    "page": element.page,
                    "route": element.route,
                    "engine": element.engine,
                    "raw_text": element.raw_text,
                    "confidence": element.confidence,
                    "bbox": element.bbox,
                    "candidate_correction": "; ".join(
                        correction.candidate_value for correction in element.candidate_corrections
                    ),
                    "correction_reason": "; ".join(
                        correction.reason for correction in element.candidate_corrections
                    ),
                    "coordinate_space": element.coordinate_space,
                }
            )
    return pd.DataFrame(records)


def table_to_dataframe(table: ReconstructedTable) -> pd.DataFrame:
    if not table.rows:
        return pd.DataFrame(columns=table.headers)

    width = max(len(table.headers), max(len(row) for row in table.rows))
    headers = normalize_headers(table.headers, width)
    normalized_rows = [row + [""] * (width - len(row)) for row in table.rows]
    return pd.DataFrame(normalized_rows, columns=headers)


def route_summary_dataframe(pages: list[PageResult]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "page": page.page,
                "route": page.route,
                "reason": page.reason,
                "native_printable_chars": page.native_printable_chars,
                "native_printable_ratio": page.native_printable_ratio,
                "image_coverage": page.image_coverage,
                "ocr_render_dpi": page.render_dpi or "",
                "text_elements": len(page.elements),
                "tables": len(page.tables),
            }
            for page in pages
        ]
    )


def render_sidebar(filename: str | None, file_bytes: bytes | None, pages: list[PageResult] | None) -> None:
    st.sidebar.subheader("Setup")
    with st.sidebar.expander("Install commands", expanded=False):
        st.code(INSTALL_COMMANDS, language="powershell")

    st.sidebar.subheader("Document metadata")
    if filename and file_bytes is not None:
        st.sidebar.write(f"Filename: `{filename}`")
        st.sidebar.write(f"Size: `{len(file_bytes) / 1024:.2f} KB`")

    if pages:
        st.sidebar.write(f"Pages: `{len(pages)}`")
        st.sidebar.write(f"Elements: `{sum(len(page.elements) for page in pages)}`")
        st.sidebar.write(f"Tables: `{sum(len(page.tables) for page in pages)}`")


def render_page_result(page: PageResult) -> None:
    metrics = st.columns(5)
    metrics[0].metric("Route", page.route)
    metrics[1].metric("Native chars", page.native_printable_chars)
    metrics[2].metric("Image coverage", f"{page.image_coverage:.2f}")
    metrics[3].metric("Text elements", len(page.elements))
    metrics[4].metric("Tables", len(page.tables))

    st.caption(f"Decision: {page.reason}")
    if page.warnings:
        for warning in page.warnings:
            st.warning(warning)

    left, right = st.columns([1, 1])
    with left:
        st.subheader("Page preview")
        if page.preview_image is not None:
            st.image(page.preview_image, use_container_width=True)
        else:
            st.info("Structured file. No page preview.")

    with right:
        st.subheader("Extracted evidence")
        df = elements_dataframe([page])
        if df.empty:
            st.info("No text elements were extracted on this page.")
        else:
            st.dataframe(df, use_container_width=True, hide_index=True)

    if page.tables:
        st.subheader("Tables and reconstructed rows")
        for table in page.tables:
            table_df = table_to_dataframe(table)
            st.caption(
                f"{table.id} | {table.engine} | {table.reconstruction_method} | "
                f"confidence {table.confidence:.2f} | coordinate space: {table.coordinate_space}"
            )
            st.dataframe(table_df, use_container_width=True, hide_index=True)
            st.download_button(
                "Download table CSV",
                data=table_df.to_csv(index=False),
                file_name=f"{table.id.lower()}.csv",
                mime="text/csv",
                key=f"download-{table.id}",
            )


def main() -> None:
    st.set_page_config(page_title="CANONIX OCR", layout="wide")
    st.title("CANONIX Document Intelligence OCR")

    st.caption(
        "Hybrid extraction for PDFs and images, plus native extraction from spreadsheets, Word, PowerPoint, and text files."
    )

    with st.sidebar:
        language_input = st.text_input("OCR languages", value="en", help="Comma-separated EasyOCR language codes.")
        languages = tuple(code.strip() for code in language_input.split(",") if code.strip()) or ("en",)
        use_gpu = st.checkbox("Use GPU if available", value=False)

    uploaded_file = st.file_uploader(
        "Upload a document, spreadsheet, image, or text file",
        type=SUPPORTED_UPLOAD_TYPES,
    )

    pages: list[PageResult] | None = None
    file_bytes: bytes | None = None
    filename: str | None = None

    if uploaded_file is not None:
        try:
            with st.spinner("Extracting document content..."):
                pages, file_bytes, filename = process_uploaded_file(uploaded_file, languages, use_gpu)
        except RuntimeError as exc:
            render_sidebar(uploaded_file.name, uploaded_file.getvalue(), None)
            st.error(str(exc))
            st.stop()
        except Exception as exc:
            render_sidebar(uploaded_file.name, uploaded_file.getvalue(), None)
            st.error(f"Extraction failed: {exc}")
            st.stop()

    render_sidebar(filename, file_bytes, pages)

    if not pages:
        st.info("Upload a PDF, Office document, spreadsheet, image, or text file to begin.")
        return

    st.subheader("Routing summary")
    st.dataframe(route_summary_dataframe(pages), use_container_width=True, hide_index=True)

    all_elements = elements_dataframe(pages)
    downloads = st.columns(2)
    downloads[0].download_button(
        "Download full extraction JSON",
        data=pages_to_json(pages),
        file_name="canonix_extraction.json",
        mime="application/json",
    )
    downloads[1].download_button(
        "Download all extracted text CSV",
        data=all_elements.to_csv(index=False),
        file_name="canonix_elements.csv",
        mime="text/csv",
        disabled=all_elements.empty,
    )

    tabs = st.tabs([f"Page {page.page}" for page in pages])
    for tab, page in zip(tabs, pages):
        with tab:
            render_page_result(page)


if __name__ == "__main__":
    main()
